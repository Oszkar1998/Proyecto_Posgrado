from flask import Blueprint, render_template, request, url_for, send_file, abort, redirect, flash, jsonify
import io
import re
import unicodedata
from datetime import datetime
from auth import permission_required
from config import direccionBase, direccionImagen, direccionPDF
from db import get_connection

actas_grado_bp = Blueprint('actas_grado', __name__, template_folder='../templates')


def _filtrar_por_nivel(plan):
    """Determina el nivel según el plan de estudios."""
    try:
        p = int(float(plan))
        if 3000 <= p <= 3999:
            return "Especialidad"
        elif 4000 <= p <= 4999:
            return "Maestría"
        elif 5000 <= p <= 5999:
            return "Doctorado"
    except (TypeError, ValueError):
        pass
    return None

def _clave_folio(folio):
    """Clave de orden natural para folios tipo 'DP-033', 'MP-172'."""
    partes = re.split(r"(\d+)", str(folio or "").upper())
    return [(0, int(p), "") if p.isdigit() else (1, 0, p) for p in partes if p != ""]


CAMPOS_BUSQUEDA = ("numDeCuenta", "folio", "nombreAlumno", "sinodal", "nivel", "fechaInicio", "fechaFin")
NIVELES = ("Especialidad", "Maestría", "Doctorado")
SIN_NIVEL = "Sin nivel identificado"


def _leer_criterios(form):
    """Toma del formulario los criterios de búsqueda (los mismos para pantalla y Excel)."""
    return {campo: (form.get(campo) or "").strip() for campo in CAMPOS_BUSQUEDA}

def _buscar_actas(criterios):
    """
    Ejecuta la búsqueda de actas con los criterios indicados.

    Devuelve (resultados, mensaje, directo):
      - resultados: lista de (folio, numCuenta, nombreCompleto), sin duplicados y ordenada por folio.
      - mensaje: aviso para el usuario o None.
      - directo: (folio, numCuenta, nombre) cuando cuenta + folio coinciden con una sola acta.

    La usan la pantalla de búsqueda y la descarga a Excel, para que ambas den lo mismo.
    """
    resultados = []
    mensaje = None

    numDeCuenta = criterios.get("numDeCuenta", "").strip()
    folio = criterios.get("folio", "").strip()
    sinodal = criterios.get("sinodal", "").strip()
    fecha_inicio = criterios.get("fechaInicio", "").strip()
    fecha_fin = criterios.get("fechaFin", "").strip()
    nombreAlumno = criterios.get("nombreAlumno", "").strip()
    nivel = criterios.get("nivel", "").strip()
    if nivel not in ("Especialidad", "Maestría", "Doctorado"):
        nivel = ""

    with get_connection() as conn:
        cursor = conn.cursor()

        if numDeCuenta and not folio and not sinodal:
            cursor.execute("SELECT folio, numCuenta, nombreCompleto FROM actasGrado WHERE numCuenta = ?", (numDeCuenta,))
            resultados = cursor.fetchall()
            if not resultados:
                mensaje = "No se encontraron resultados para el número de cuenta proporcionado."

        elif folio and not numDeCuenta and not sinodal:
            cursor.execute("SELECT folio, numCuenta, nombreCompleto FROM actasGrado WHERE folio = ?", (folio,))
            resultados = cursor.fetchall()
            if not resultados:
                mensaje = "No se encontraron resultados para el folio proporcionado."

        elif folio and numDeCuenta and not sinodal:
            cursor.execute("SELECT folio, numCuenta, nombreCompleto FROM actasGrado WHERE folio = ? AND numCuenta = ?", (folio, numDeCuenta))
            resultado = cursor.fetchone()
            if resultado:
                # Coincidencia exacta: la pantalla abre el PDF directamente
                return [resultado], None, resultado
            else:
                mensaje = "No se encontró ninguna acta con ese folio y número de cuenta."

        if sinodal:
            nombre_buscar = sinodal.lower()
            rango_fechas_valido = False
            f_ini = f_fin = None
            if fecha_inicio and fecha_fin:
                try:
                    f_ini = datetime.strptime(fecha_inicio, "%Y-%m-%d")
                    f_fin = datetime.strptime(fecha_fin, "%Y-%m-%d")
                    rango_fechas_valido = True
                except ValueError:
                    mensaje = "Formato de fecha no válido. Usa el calendario."

            cursor.execute("SELECT * FROM actasGrado")
            filas = cursor.fetchall()
            columnas = [d[0] for d in cursor.description]

            try:
                i_folio = columnas.index("folio")
                i_cuenta = columnas.index("numCuenta")
                i_nombre = columnas.index("nombreCompleto")
            except ValueError:
                return [], "Estructura de la base no esperada (faltan columnas básicas).", None

            claves = ["sinodal", "presidente", "secretario", "vocal", "jurado", "asesor", "director", "tutor"]
            idx_jurado = [i for i, c in enumerate(columnas) if any(k in c.lower() for k in claves)]

            i_anio = columnas.index("anio") if "anio" in columnas else None
            i_mes  = columnas.index("mes")  if "mes"  in columnas else None
            i_dia  = columnas.index("dia")  if "dia"  in columnas else None

            for fila in filas:
                if not idx_jurado:
                    continue
                textos_jurado = [str(fila[i] or "").lower() for i in idx_jurado]
                if any(nombre_buscar in t for t in textos_jurado):
                    if rango_fechas_valido and (i_anio is not None and i_mes is not None and i_dia is not None):
                        try:
                            f_acta = datetime(int(fila[i_anio]), int(fila[i_mes]), int(fila[i_dia]))
                            if not (f_ini <= f_acta <= f_fin):
                                continue
                        except Exception:
                            continue


                    resultados.append((fila[i_folio], fila[i_cuenta], fila[i_nombre]))

            if not resultados and not mensaje:
                mensaje = "No se encontraron actas con ese sinodal."

        if not sinodal and not nombreAlumno and fecha_inicio and fecha_fin:
            try:
                f_ini = datetime.strptime(fecha_inicio, "%Y-%m-%d")
                f_fin = datetime.strptime(fecha_fin, "%Y-%m-%d")

                cursor.execute("""
                    SELECT folio, numCuenta, nombreCompleto, anio, mes, dia
                    FROM actasGrado
                """)
                todas = cursor.fetchall()
                for folio_x, cuenta_x, nombre_x, anio_x, mes_x, dia_x in todas:
                    try:
                        f_acta = datetime(int(anio_x), int(mes_x), int(dia_x))
                        if f_ini <= f_acta <= f_fin:
                            resultados.append((folio_x, cuenta_x, nombre_x))
                    except Exception:
                        continue

                if not resultados:
                    mensaje = "No se encontraron resultados en el rango de fechas proporcionado."
            except ValueError:
                mensaje = "Formato de fecha no válido. Usa el calendario."

        if nombreAlumno and not (folio and numDeCuenta):
            nombre_buscar = nombreAlumno.strip().upper()
            rango_fechas_valido = False
            f_ini = f_fin = None
            if fecha_inicio and fecha_fin:
                try:
                    f_ini = datetime.strptime(fecha_inicio, "%Y-%m-%d")
                    f_fin = datetime.strptime(fecha_fin, "%Y-%m-%d")
                    rango_fechas_valido = True
                except ValueError:
                    mensaje = "Formato de fecha no válido. Usa el calendario."

            cursor.execute("SELECT folio, numCuenta, nombreCompleto, anio, mes, dia FROM actasGrado")
            filas = cursor.fetchall()

            for folio_x, cuenta_x, nombre_x, anio_x, mes_x, dia_x in filas:
                if nombre_x and (nombre_buscar in nombre_x.upper()):
                    if rango_fechas_valido:
                        try:
                            f_acta = datetime(int(anio_x), int(mes_x), int(dia_x))
                            if not (f_ini <= f_acta <= f_fin):
                                continue
                        except Exception:
                            continue
                    resultados.append((folio_x, cuenta_x, nombre_x))

            if not resultados and not mensaje:
                mensaje = "No se encontraron actas con ese nombre."

        # ---------- Búsqueda / filtro por NIVEL (según plan de estudios) ----------
        if nivel:
            otros_criterios = numDeCuenta or folio or sinodal or nombreAlumno or (fecha_inicio and fecha_fin)

            if not otros_criterios:
                # Solo se eligió el nivel: traer todas las actas de ese nivel
                cursor.execute("SELECT folio, numCuenta, nombreCompleto, planEstudios FROM actasGrado")
                for folio_x, cuenta_x, nombre_x, plan_x in cursor.fetchall():
                    if _filtrar_por_nivel(plan_x) == nivel:
                        resultados.append((folio_x, cuenta_x, nombre_x))
            elif resultados:
                # Hay otros criterios: dejar solo las actas de ese nivel
                cursor.execute("SELECT folio, numCuenta, planEstudios FROM actasGrado")
                planes = {(str(f), str(c)): p for f, c, p in cursor.fetchall()}
                resultados = [
                    r for r in resultados
                    if _filtrar_por_nivel(planes.get((str(r[0]), str(r[1])))) == nivel
                ]

            if not resultados and (not mensaje or mensaje.startswith("No se encontr")):
                mensaje = f"No se encontraron actas de {nivel} con los criterios indicados."

        # Quitar duplicados conservando el orden
        vistos = set()
        unicos = []
        for r in resultados:
            clave = (str(r[0]), str(r[1]))
            if clave not in vistos:
                vistos.add(clave)
                unicos.append(r)
        resultados = unicos

        # Ordenar por folio de forma ascendente (orden natural: DP-2 antes que DP-10)
        resultados.sort(key=lambda r: _clave_folio(r[0]))

    return resultados, mensaje, None

def _texto_criterios(criterios):
    """Resumen legible de los criterios usados, para el encabezado del Excel."""
    partes = []
    if criterios.get("numDeCuenta"):
        partes.append(f"Cuenta: {criterios['numDeCuenta']}")
    if criterios.get("folio"):
        partes.append(f"Folio: {criterios['folio']}")
    if criterios.get("nombreAlumno"):
        partes.append(f"Nombre: {criterios['nombreAlumno'].upper()}")
    if criterios.get("sinodal"):
        partes.append(f"Sinodal: {criterios['sinodal'].upper()}")
    if criterios.get("fechaInicio") and criterios.get("fechaFin"):
        try:
            ini = datetime.strptime(criterios["fechaInicio"], "%Y-%m-%d").strftime("%d/%m/%Y")
            fin = datetime.strptime(criterios["fechaFin"], "%Y-%m-%d").strftime("%d/%m/%Y")
            partes.append(f"Fechas: {ini} a {fin}")
        except ValueError:
            pass
    nivel = criterios.get("nivel")
    partes.append(f"Nivel: {nivel if nivel in NIVELES else 'Todos'}")
    return " · ".join(partes)

def _plan_legible(plan):
    """4170.0 -> 4170 (planEstudios está guardado como REAL)."""
    if plan is None or str(plan).strip() == "":
        return ""
    try:
        numero = float(plan)
        return int(numero) if numero == int(numero) else numero
    except (TypeError, ValueError):
        return str(plan).strip()

def _clave_nombre(nombre):
    """Orden alfabético en español: ignora acentos y mayúsculas, y la Ñ va después de la N."""
    texto = str(nombre or "").upper().replace("Ñ", "\x00")
    texto = unicodedata.normalize("NFD", texto)
    texto = "".join(c for c in texto if unicodedata.category(c) != "Mn")
    return texto.replace("\x00", "N\U0010FFFF")

def _actas(n):
    return f"{n} acta" if n == 1 else f"{n} actas"

def _generar_excel_actas(grupos, criterios):
    """
    Arma el Excel de resultados en memoria.

    grupos: lista de (nivel, filas) en el orden en que deben aparecer; cada fila es
            (numCuenta, nombreCompleto, planEstudios, semestreIngreso, semestreEgreso).
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    AZUL, AZUL_CLARO, AZUL_TENUE = "0C447C", "E6F1FB", "B5D4F4"
    ORO, GRIS, CEBRA = "F7C948", "F1EFE8", "F7F9FC"

    def relleno(color):
        return PatternFill("solid", start_color=color, end_color=color)

    linea = Side(style="thin", color="D3D1C7")
    borde = Border(left=linea, right=linea, top=linea, bottom=linea)
    centro = Alignment(horizontal="center", vertical="center")
    izquierda = Alignment(horizontal="left", vertical="center", indent=1)

    wb = Workbook()
    ws = wb.active
    ws.title = "Actas de grado"
    ws.sheet_view.showGridLines = False

    COLUMNAS = ["No.", "Núm. de cuenta", "Nombre completo", "Plan de estudios", "Ingreso", "Egreso"]
    ULTIMA = len(COLUMNAS)

    def franja(fila, texto, fuente, color, alineacion=centro, alto=20):
        """Escribe un texto que abarca todas las columnas de la fila."""
        ws.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=ULTIMA)
        celda = ws.cell(row=fila, column=1, value=texto)
        celda.font = fuente
        celda.alignment = alineacion
        if color:
            for col in range(1, ULTIMA + 1):
                ws.cell(row=fila, column=col).fill = relleno(color)
        ws.row_dimensions[fila].height = alto

    total = sum(len(filas) for _, filas in grupos)
    conteo = {nivel: len(filas) for nivel, filas in grupos}
    desglose = [f"{nivel} {conteo.get(nivel, 0)}" for nivel in NIVELES]
    if conteo.get(SIN_NIVEL):
        desglose.append(f"{SIN_NIVEL} {conteo[SIN_NIVEL]}")

    # ---------- Encabezado del archivo ----------
    franja(1, "Coordinación de Posgrado · Facultad de Ingeniería, UNAM",
           Font(name="Arial", size=14, bold=True, color="FFFFFF"), AZUL, alto=28)
    franja(2, "Actas de grado · Resultados de búsqueda",
           Font(name="Arial", size=11, color=AZUL_TENUE), AZUL, alto=20)
    franja(3, f"Criterios: {_texto_criterios(criterios)}",
           Font(name="Arial", size=10, color="444441"), None, izquierda, alto=18)
    franja(4, f"Total: {_actas(total)} ({' · '.join(desglose)})  ·  "
              f"Generado: {datetime.now().strftime('%d/%m/%Y %H:%M')}",
           Font(name="Arial", size=10, color="444441"), None, izquierda, alto=18)

    # ---------- Un listado por nivel ----------
    fila = 6
    largo_nombre = len(COLUMNAS[2])
    for nivel, filas in grupos:
        if not filas:
            continue

        franja(fila, f"{nivel} ({len(filas)})",
               Font(name="Arial", size=12, bold=True, color="412402"), ORO, izquierda, alto=22)
        fila += 1

        for col, titulo in enumerate(COLUMNAS, start=1):
            celda = ws.cell(row=fila, column=col, value=titulo)
            celda.font = Font(name="Arial", size=10, bold=True, color=AZUL)
            celda.fill = relleno(AZUL_CLARO)
            celda.alignment = centro
            celda.border = borde
        ws.row_dimensions[fila].height = 20
        fila += 1

        for n, (cuenta, nombre, plan, ingreso, egreso) in enumerate(filas, start=1):
            nombre = str(nombre or "").strip()
            largo_nombre = max(largo_nombre, len(nombre))
            valores = [n, cuenta, nombre, _plan_legible(plan),
                       str(ingreso or "").strip(), str(egreso or "").strip()]
            for col, valor in enumerate(valores, start=1):
                celda = ws.cell(row=fila, column=col, value=valor)
                celda.font = Font(name="Arial", size=10, color="2C2C2A")
                celda.alignment = izquierda if col == 3 else centro
                celda.border = borde
                if n % 2 == 0:
                    celda.fill = relleno(CEBRA)
                if col in (2, 4) and isinstance(valor, (int, float)):
                    celda.number_format = "0"      # sin separador de miles
            ws.row_dimensions[fila].height = 18
            fila += 1

        fila += 1   # renglón en blanco entre niveles

    franja(fila, f"Total general: {_actas(total)}",
           Font(name="Arial", size=11, bold=True, color="444441"), GRIS, izquierda, alto=22)

    # ---------- Anchos e impresión ----------
    anchos = [6, 17, min(max(largo_nombre + 4, 34), 60), 17, 12, 12]
    for letra, ancho in zip("ABCDEF", anchos):
        ws.column_dimensions[letra].width = ancho

    ws.page_setup.paperSize = ws.PAPERSIZE_LETTER
    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = ws.page_margins.right = 0.5
    ws.page_margins.top = ws.page_margins.bottom = 0.6
    ws.oddFooter.center.text = "Página &P de &N"
    ws.oddFooter.center.size = 9

    salida = io.BytesIO()
    wb.save(salida)
    salida.seek(0)
    return salida


#################################################################################################################################
@actas_grado_bp.route("/", methods=["GET", "POST"])
@permission_required("actas_grado")
def index():

    resultados = []
    mensaje = None
    criterios = {}

    if request.method == "POST":
        criterios = _leer_criterios(request.form)
        resultados, mensaje, directo = _buscar_actas(criterios)

        if directo:
            pdf_url = url_for('actas_grado.ver_pdf', num_cuenta=directo[1], folio=directo[0])
            return f"""
                <script>
                    window.open('{pdf_url}', '_blank');
                    window.location.href = '{url_for('actas_grado.index')}';
                </script>
            """

    return render_template("actasGrado.html", resultados=resultados, mensaje=mensaje, criterios=criterios)

@actas_grado_bp.route("/exportar_excel", methods=["POST"])
@permission_required("actas_grado")
def exportar_excel():
    """Descarga en Excel los mismos resultados de la búsqueda, separados por nivel."""
    criterios = _leer_criterios(request.form)
    resultados, _mensaje, _directo = _buscar_actas(criterios)

    if not resultados:
        flash("❌ No hay resultados para descargar. Realiza primero una búsqueda.", "error")
        return redirect(url_for("actas_grado.index"))

    # Datos que van en el Excel y que la lista en pantalla no trae
    with get_connection() as conn:
        cursor = conn.cursor()
        cursor.execute("""
            SELECT folio, numCuenta, planEstudios, semestreIngreso, semestreEgreso
            FROM actasGrado
        """)
        detalle = {(str(f), str(c)): (plan, ingreso, egreso)
                   for f, c, plan, ingreso, egreso in cursor.fetchall()}

    por_nivel = {nivel: [] for nivel in NIVELES + (SIN_NIVEL,)}
    for folio, cuenta, nombre in resultados:
        plan, ingreso, egreso = detalle.get((str(folio), str(cuenta)), (None, "", ""))
        nivel = _filtrar_por_nivel(plan) or SIN_NIVEL
        por_nivel[nivel].append((cuenta, nombre, plan, ingreso, egreso))

    grupos = []
    for nivel in NIVELES + (SIN_NIVEL,):
        filas = sorted(por_nivel[nivel], key=lambda f: (_clave_nombre(f[1]), str(f[0])))
        grupos.append((nivel, filas))

    try:
        archivo = _generar_excel_actas(grupos, criterios)
    except ImportError:
        flash("❌ Falta instalar la librería openpyxl en el servidor (pip install openpyxl).", "error")
        return redirect(url_for("actas_grado.index"))

    nombre_archivo = f"actas_grado_{datetime.now().strftime('%Y-%m-%d_%H%M')}.xlsx"
    return send_file(
        archivo,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name=nombre_archivo,
    )

@actas_grado_bp.route("/ver_pdf/<num_cuenta>/<folio>")
def ver_pdf(num_cuenta, folio):
    with get_connection() as conn:
        cursor = conn.cursor()
        cursor.execute("SELECT actaPDF FROM actasGrado WHERE numCuenta=? AND folio=?", (num_cuenta, folio))
        archivo = cursor.fetchone()

    if archivo and archivo[0]:
        return send_file(io.BytesIO(archivo[0]), mimetype='application/pdf')
    else:
        return abort(404)

@actas_grado_bp.after_request
def disable_cache(response):
    response.headers["Cache-Control"] = "no-cache, no-store, must-revalidate"
    response.headers["Pragma"] = "no-cache"
    response.headers["Expires"] = "0"
    return response

@actas_grado_bp.route("/guardar", methods=["POST"])
def guardar_acta():
    folio = (request.form.get("folio") or "").strip()
    numCuenta = (request.form.get("numCuenta") or "").strip()
    nombreCompleto = (request.form.get("nombreCompleto") or "").strip()
    grado = (request.form.get("grado") or "").strip()
    semestreIngreso = (request.form.get("semestreIngreso") or "").strip()
    semestreEgreso = (request.form.get("semestreEgreso") or "").strip()
    promedio = (request.form.get("promedio") or "").strip()
    planEstudios = (request.form.get("planEstudios") or "").strip()
    dia = (request.form.get("dia") or "").strip()
    mes = (request.form.get("mes") or "").strip()
    anio = (request.form.get("anio") or "").strip()
    nacionalidad = (request.form.get("nacionalidad") or "").strip()
    resultado = (request.form.get("resultado") or "").strip()
    presidente = (request.form.get("presidente") or "").strip()
    secretario = (request.form.get("secretario") or "").strip()
    vocal1 = (request.form.get("vocal1") or "").strip()
    vocal2 = (request.form.get("vocal2") or "").strip()
    vocal3 = (request.form.get("vocal3") or "").strip()
    sexo = (request.form.get("sexo") or "").strip()

    required = [folio, numCuenta, nombreCompleto, grado, semestreIngreso, semestreEgreso,
                dia, mes, anio, resultado, presidente, secretario, vocal1, sexo]
    if any(not x for x in required):
        return render_template("actasGrado.html", resultados=[], mensaje="Faltan campos obligatorios (*) para guardar.")

    f = request.files.get("actaPDF")
    if not f or f.filename == "":
        return render_template("actasGrado.html", resultados=[], mensaje="Debes adjuntar el PDF del acta.")
    pdf_bytes = f.read()

    if not pdf_bytes.startswith(b"%PDF"):
        return render_template("actasGrado.html", resultados=[], mensaje="El archivo adjunto no parece ser un PDF válido.")

    if len(pdf_bytes) > 15 * 1024 * 1024:
        return render_template("actasGrado.html", resultados=[], mensaje="El PDF excede 15MB. Comprime o reduce el archivo.")

    try:
        numCuenta_int = int(numCuenta)
        dia_int = int(dia); mes_int = int(mes); anio_int = int(anio)
    except ValueError:
        return render_template("actasGrado.html", resultados=[], mensaje="NumCuenta/Día/Mes/Año deben ser numéricos.")

    prom_val = None
    if promedio != "":
        try:
            prom_val = float(promedio)
        except ValueError:
            return render_template("actasGrado.html", resultados=[], mensaje="Promedio debe ser numérico (ej: 9.3).")

    plan_val = None
    if planEstudios != "":
        try:
            plan_val = float(planEstudios)
        except ValueError:
            return render_template("actasGrado.html", resultados=[], mensaje="Plan de estudios debe ser numérico (según tu BD).")

    with get_connection() as conn:
        cursor = conn.cursor()

        cursor.execute("SELECT 1 FROM actasGrado WHERE numCuenta=? AND folio=? LIMIT 1", (numCuenta_int, folio))
        if cursor.fetchone():
            return render_template("actasGrado.html", resultados=[], mensaje="Ya existe un registro con ese Folio y Número de Cuenta (duplicado).")

        cursor.execute("""
            SELECT 1 FROM actasGrado
            WHERE nombreCompleto=? AND anio=? AND mes=? AND dia=? AND presidente=? AND secretario=? AND vocal1=? LIMIT 1
        """, (nombreCompleto, anio_int, mes_int, dia_int, presidente, secretario, vocal1))
        if cursor.fetchone():
            return render_template("actasGrado.html", resultados=[], mensaje="Parece duplicado: ya existe un acta con mismo alumno/fecha/jurado.")

        cursor.execute("""
            INSERT INTO actasGrado (
              folio, numCuenta, nombreCompleto, grado, semestreIngreso, semestreEgreso,
              promedio, planEstudios, dia, mes, anio, nacionalidad, resultado,
              presidente, secretario, vocal1, vocal2, vocal3, sexo, actaPDF
            )
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """, (
            folio, numCuenta_int, nombreCompleto, grado, semestreIngreso, semestreEgreso,
            prom_val, plan_val, dia_int, mes_int, anio_int, nacionalidad, resultado,
            presidente, secretario, vocal1, vocal2, vocal3, sexo, bytes(pdf_bytes)
        ))
        conn.commit()

    flash("✅ Se guardaron correctamente los datos del acta.", "success")
    return redirect(url_for("actas_grado.index"))

@actas_grado_bp.route("/get_acta", methods=["POST"])
def get_acta_json():
    data = request.get_json(silent=True) or {}
    numCuenta = (data.get("numCuenta") or "").strip()
    folio = (data.get("folio") or "").strip()

    if not numCuenta or not folio:
        return {"error": "Faltan numCuenta y folio."}, 400

    with get_connection() as conn:
        cursor = conn.cursor()
        cursor.execute("""
            SELECT folio, numCuenta, nombreCompleto, grado, semestreIngreso, semestreEgreso,
                   promedio, planEstudios, dia, mes, anio, nacionalidad, resultado,
                   presidente, secretario, vocal1, vocal2, vocal3, sexo
            FROM actasGrado
            WHERE numCuenta=? AND folio=?
            LIMIT 1
        """, (numCuenta, folio))
        row = cursor.fetchone()

    if not row:
        return {"error": "Acta no encontrada."}, 404

    keys = ["folio","numCuenta","nombreCompleto","grado","semestreIngreso","semestreEgreso",
            "promedio","planEstudios","dia","mes","anio","nacionalidad","resultado",
            "presidente","secretario","vocal1","vocal2","vocal3","sexo"]
    return dict(zip(keys, row)), 200

@actas_grado_bp.route("/actualizar", methods=["POST"])
def actualizar_acta():
    old_num = (request.form.get("old_numCuenta") or "").strip()
    old_fol = (request.form.get("old_folio") or "").strip()

    folio = (request.form.get("folio") or "").strip()
    numCuenta = (request.form.get("numCuenta") or "").strip()
    nombreCompleto = (request.form.get("nombreCompleto") or "").strip()
    grado = (request.form.get("grado") or "").strip()
    semestreIngreso = (request.form.get("semestreIngreso") or "").strip()
    semestreEgreso = (request.form.get("semestreEgreso") or "").strip()
    promedio = (request.form.get("promedio") or "").strip()
    planEstudios = (request.form.get("planEstudios") or "").strip()
    dia = (request.form.get("dia") or "").strip()
    mes = (request.form.get("mes") or "").strip()
    anio = (request.form.get("anio") or "").strip()
    nacionalidad = (request.form.get("nacionalidad") or "").strip()
    resultado = (request.form.get("resultado") or "").strip()
    presidente = (request.form.get("presidente") or "").strip()
    secretario = (request.form.get("secretario") or "").strip()
    vocal1 = (request.form.get("vocal1") or "").strip()
    vocal2 = (request.form.get("vocal2") or "").strip()
    vocal3 = (request.form.get("vocal3") or "").strip()
    sexo = (request.form.get("sexo") or "").strip()

    required = [old_num, old_fol, folio, numCuenta, nombreCompleto, grado, semestreIngreso, semestreEgreso,
                dia, mes, anio, resultado, presidente, secretario, vocal1, sexo]
    if any(not x for x in required):
        flash("❌ Faltan campos obligatorios para actualizar.", "error")
        return redirect(url_for("actas_grado.index", open="modificar"))

    try:
        numCuenta_int = int(numCuenta)
        old_num_int = int(old_num)
        dia_int = int(dia); mes_int = int(mes); anio_int = int(anio)
    except ValueError:
        flash("❌ NumCuenta/Día/Mes/Año deben ser numéricos.", "error")
        return redirect(url_for("actas_grado.index", open="modificar"))

    prom_val = None
    if promedio != "":
        try: prom_val = float(promedio)
        except ValueError:
            flash("❌ Promedio inválido.", "error")
            return redirect(url_for("actas_grado.index", open="modificar"))

    plan_val = None
    if planEstudios != "":
        try: plan_val = float(planEstudios)
        except ValueError:
            flash("❌ Plan de estudios inválido.", "error")
            return redirect(url_for("actas_grado.index", open="modificar"))

    new_pdf_bytes = None
    f = request.files.get("actaPDF")
    if f and f.filename:
        new_pdf_bytes = f.read()
        if not new_pdf_bytes.startswith(b"%PDF"):
            flash("❌ El archivo adjunto no es PDF válido.", "error")
            return redirect(url_for("actas_grado.index", open="modificar"))

    with get_connection() as conn:
        cursor = conn.cursor()

        cursor.execute("SELECT 1 FROM actasGrado WHERE numCuenta=? AND folio=? LIMIT 1", (old_num_int, old_fol))
        if not cursor.fetchone():
            flash("❌ No se encontró el registro original para actualizar.", "error")
            return redirect(url_for("actas_grado.index", open="modificar"))

        if (old_num_int != numCuenta_int) or (old_fol != folio):
            cursor.execute("SELECT 1 FROM actasGrado WHERE numCuenta=? AND folio=? LIMIT 1", (numCuenta_int, folio))
            if cursor.fetchone():
                flash("❌ Ya existe otro registro con ese Folio y Número de Cuenta.", "error")
                return redirect(url_for("actas_grado.index", open="modificar"))

        if new_pdf_bytes is None:
            cursor.execute("""
                UPDATE actasGrado
                SET folio=?, numCuenta=?, nombreCompleto=?, grado=?, semestreIngreso=?, semestreEgreso=?,
                    promedio=?, planEstudios=?, dia=?, mes=?, anio=?, nacionalidad=?, resultado=?,
                    presidente=?, secretario=?, vocal1=?, vocal2=?, vocal3=?, sexo=?
                WHERE numCuenta=? AND folio=?
            """, (
                folio, numCuenta_int, nombreCompleto, grado, semestreIngreso, semestreEgreso,
                prom_val, plan_val, dia_int, mes_int, anio_int, nacionalidad, resultado,
                presidente, secretario, vocal1, vocal2, vocal3, sexo,
                old_num_int, old_fol
            ))
        else:
            cursor.execute("""
                UPDATE actasGrado
                SET folio=?, numCuenta=?, nombreCompleto=?, grado=?, semestreIngreso=?, semestreEgreso=?,
                    promedio=?, planEstudios=?, dia=?, mes=?, anio=?, nacionalidad=?, resultado=?,
                    presidente=?, secretario=?, vocal1=?, vocal2=?, vocal3=?, sexo=?, actaPDF=?
                WHERE numCuenta=? AND folio=?
            """, (
                folio, numCuenta_int, nombreCompleto, grado, semestreIngreso, semestreEgreso,
                prom_val, plan_val, dia_int, mes_int, anio_int, nacionalidad, resultado,
                presidente, secretario, vocal1, vocal2, vocal3, sexo, bytes(new_pdf_bytes),
                old_num_int, old_fol
            ))
        conn.commit()

    flash("✅ Cambios guardados correctamente.", "success")
    return redirect(url_for("actas_grado.index"))

@actas_grado_bp.route("/guardar_fecha_extra", methods=["POST"])
def guardar_fecha_extra():
    numCuenta = (request.form.get("numCuenta") or "").strip()
    folio = (request.form.get("folio") or "").strip()
    fecha = (request.form.get("fecha") or "").strip()
    campo = (request.form.get("campo") or "").strip()

    if not numCuenta or not folio or not fecha or not campo:
        flash("❌ Debes ingresar número de cuenta, folio y fecha.", "error")
        return redirect(url_for("actas_grado.index"))

    if campo not in ["pagoSinodales", "dgaeEntregaCarpeta"]:
        flash("❌ Campo no válido.", "error")
        return redirect(url_for("actas_grado.index"))

    try:
        numCuenta_int = int(numCuenta)
    except ValueError:
        flash("❌ El número de cuenta debe ser numérico.", "error")
        return redirect(url_for("actas_grado.index", open=campo))

    with get_connection() as conn:
        cursor = conn.cursor()

        cursor.execute("""
            SELECT nombreCompleto FROM actasGrado
            WHERE numCuenta = ? AND folio = ? LIMIT 1
        """, (numCuenta_int, folio))
        acta = cursor.fetchone()

        if not acta:
            flash("❌ No se encontró ningún registro con ese número de cuenta y folio.", "error")
            return redirect(url_for("actas_grado.index", open=campo))

        if campo == "pagoSinodales":
            cursor.execute("""
                UPDATE actasGrado SET pagoSinodales = ?
                WHERE numCuenta = ? AND folio = ?
            """, (fecha, numCuenta_int, folio))
            mensaje = "✅ Fecha de pago de sinodales guardada correctamente."
        else:
            cursor.execute("""
                UPDATE actasGrado SET dgaeEntregaCarpeta = ?
                WHERE numCuenta = ? AND folio = ?
            """, (fecha, numCuenta_int, folio))
            mensaje = "✅ Fecha de entrega de carpeta DGAE guardada correctamente."

        conn.commit()

    flash(mensaje, "success")
    return redirect(url_for("actas_grado.index"))

@actas_grado_bp.route("/buscar_pago_sinodales", methods=["POST"])
def buscar_pago_sinodales():
    data = request.get_json(silent=True) or {}

    sinodal = (data.get("sinodal") or "").strip().lower()
    fecha_inicio = (data.get("fechaInicio") or "").strip()
    fecha_fin = (data.get("fechaFin") or "").strip()

    if not sinodal:
        return jsonify({"ok": False, "error": "Ingresa el nombre del sinodal."}), 400

    f_ini = f_fin = None
    if fecha_inicio and fecha_fin:
        try:
            f_ini = datetime.strptime(fecha_inicio, "%Y-%m-%d")
            f_fin = datetime.strptime(fecha_fin, "%Y-%m-%d")
        except ValueError:
            return jsonify({"ok": False, "error": "Formato de fecha inválido."}), 400

    with get_connection() as conn:
        conn.row_factory = __import__('sqlite3').Row
        cursor = conn.cursor()
        cursor.execute("""
            SELECT folio, numCuenta, nombreCompleto, dia, mes, anio,
                   presidente, secretario, vocal1, vocal2, vocal3,
                   pagoSinodales, dgaeEntregaCarpeta
            FROM actasGrado
        """)
        filas = cursor.fetchall()

    resultados = []
    for fila in filas:
        jurado = [fila["presidente"], fila["secretario"], fila["vocal1"], fila["vocal2"], fila["vocal3"]]
        coincide_sinodal = any(sinodal in str(nombre or "").lower() for nombre in jurado)

        if not coincide_sinodal:
            continue

        if f_ini and f_fin:
            try:
                fecha_acta = datetime(int(fila["anio"]), int(fila["mes"]), int(fila["dia"]))
                if not (f_ini <= fecha_acta <= f_fin):
                    continue
            except Exception:
                continue

        if not fila["pagoSinodales"]:
            continue

        resultados.append({
            "folio": fila["folio"],
            "numCuenta": fila["numCuenta"],
            "nombreCompleto": fila["nombreCompleto"],
            "fechaActa": f'{fila["dia"]}/{fila["mes"]}/{fila["anio"]}',
            "pagoSinodales": fila["pagoSinodales"],
            "dgaeEntregaCarpeta": fila["dgaeEntregaCarpeta"],
        })

    return jsonify({"ok": True, "resultados": resultados})
